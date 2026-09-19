#!/usr/bin/env python3
"""Build a .pptx test deck that exercises parser/renderer features the real
sample decks don't: embedded video, animated GIF, click-to-reveal builds,
scaled group shapes, cropped + hyperlinked images, fills/theme colors, and
placeholder-inherited positions.

Dev-only tool (not shipped with the app). Needs python-pptx:
    python3 -m venv .venv && .venv/bin/pip install python-pptx
    .venv/bin/python tools/build_media_test_deck.py \
        --video clip.mp4 --poster poster.jpg --gif anim.gif \
        --image photo.jpg --out presentations_testing/media_test_deck.pptx
"""
import argparse
import copy

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

BLANK = 6
TITLE_ONLY = 5
TITLE_SLIDE = 0


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(36)
    run.font.bold = True
    return box


def click_build_timing(targets):
    """<p:timing> revealing each target on successive clicks (PowerPoint's
    "Appear" entrance effect). targets: list of (spid, paragraph_index or
    None for the whole shape)."""
    ids = iter(range(3, 1000))
    steps = []
    for spid, para in targets:
        a, b, c, d = next(ids), next(ids), next(ids), next(ids)
        txel = f'<p:txEl><p:pRg st="{para}" end="{para}"/></p:txEl>' if para is not None else ""
        steps.append(f"""
<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
 <p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
  <p:par><p:cTn id="{c}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">
   <p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
    <p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
     <p:tgtEl><p:spTgt spid="{spid}">{txel}</p:spTgt></p:tgtEl>
     <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>
     <p:to><p:strVal val="visible"/></p:to></p:set>
   </p:childTnLst></p:cTn></p:par>
 </p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>""")
    text_spids = sorted({spid for spid, para in targets if para is not None})
    bld = "".join(f'<p:bldP spid="{s}" grpId="0" build="p"/>' for s in text_spids)
    xml = f"""<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par>
<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
{''.join(steps)}
</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>
<p:bldLst>{bld}</p:bldLst></p:timing>"""
    return etree.fromstring(xml)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--video", required=True)
    ap.add_argument("--poster", required=True)
    ap.add_argument("--gif", required=True)
    ap.add_argument("--image", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    prs = Presentation()  # default 10in x 7.5in (4:3)

    # 1. Title slide: placeholders with no explicit <a:xfrm> on the slide -
    # positions inherit from the layout, exercising the parser's fallback.
    s = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE])
    s.shapes.title.text = "PlatformPresenter media test"
    s.placeholders[1].text = "Video - GIF - click builds - groups - crops - links"

    # 2. Embedded MP4 video with poster frame.
    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_title(s, "Embedded video")
    s.shapes.add_movie(args.video, Inches(1), Inches(1.5), Inches(8), Inches(8 * 720 / 1720),
                       poster_frame_image=args.poster, mime_type="video/mp4")
    cap = s.shapes.add_textbox(Inches(1), Inches(5.3), Inches(8), Inches(0.6))
    cap.text_frame.text = "Press E / X to play"

    # 3. Animated GIF.
    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_title(s, "Animated GIF")
    s.shapes.add_picture(args.gif, Inches(2), Inches(1.6), width=Inches(6))

    # 4. Click-to-reveal: bullets one by one, then an image.
    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_title(s, "Click to reveal")
    body = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5), Inches(3.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(["First point", "Second point", "Third point", "Fourth point"]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = "• " + line
        para.runs[0].font.size = Pt(28)
    pic = s.shapes.add_picture(args.image, Inches(6.2), Inches(2), width=Inches(3.3))
    targets = [(body.shape_id, i) for i in range(4)] + [(pic.shape_id, None)]
    s._element.append(click_build_timing(targets))

    # 5. Scaled group: children authored in child space, group stretched
    # 1.5x, so parsed positions must go through the group transform.
    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_title(s, "Grouped shapes (scaled 1.5x)")
    grp = s.shapes.add_group_shape()
    r = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1))
    r.fill.solid()
    r.fill.fore_color.rgb = RGBColor(0x2E, 0x86, 0xDE)
    t = grp.shapes.add_textbox(Inches(1), Inches(3.2), Inches(3), Inches(0.8))
    t.text_frame.text = "Text inside the group"
    grp.shapes.add_picture(args.image, Inches(4.2), Inches(2), width=Inches(1.4))
    xfrm = grp._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    ch_ext = xfrm.find(qn("a:chExt"))
    ext = xfrm.find(qn("a:ext"))
    ext.set("cx", str(int(int(ch_ext.get("cx")) * 1.5)))
    ext.set("cy", str(int(int(ch_ext.get("cy")) * 1.5)))
    off = xfrm.find(qn("a:off"))
    ch_off = xfrm.find(qn("a:chOff"))
    print("group: off=(%.2f,%.2f)cm chOff=(%.2f,%.2f)cm scale=1.5" % (
        Emu(int(off.get("x"))).cm, Emu(int(off.get("y"))).cm,
        Emu(int(ch_off.get("x"))).cm, Emu(int(ch_off.get("y"))).cm))
    for child in (r, t):
        ex = Emu(int(off.get("x")) + (child.left - int(ch_off.get("x"))) * 1.5).cm
        ey = Emu(int(off.get("y")) + (child.top - int(ch_off.get("y"))) * 1.5).cm
        print("  expected %-10s at (%.2f, %.2f)cm size %.2fx%.2fcm" % (
            type(child).__name__, ex, ey, Emu(child.width).cm * 1.5, Emu(child.height).cm * 1.5))

    # 6. Cropped image with a hyperlink, plus a hyperlinked text run.
    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_title(s, "Cropped image + links")
    pic = s.shapes.add_picture(args.image, Inches(1), Inches(1.6), width=Inches(5))
    pic.crop_left, pic.crop_right, pic.crop_top = 0.25, 0.1, 0.15
    pic.click_action.hyperlink.address = "https://godotengine.org"
    tb = s.shapes.add_textbox(Inches(6.3), Inches(2), Inches(3.2), Inches(1.5))
    tb.text_frame.word_wrap = True
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Linked text (Godot docs)"
    run.hyperlink.address = "https://docs.godotengine.org"

    # 7. Fills, theme colors, and rich text styling.
    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_title(s, "Shapes, colors, text styles")
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4), Inches(2.2))
    box.fill.solid()
    box.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for text, kw in [("Bold ", dict(bold=True)), ("italic ", dict(italic=True)), ("big", dict(size=Pt(40)))]:
        rr = p.add_run()
        rr.text = text
        rr.font.bold = kw.get("bold")
        rr.font.italic = kw.get("italic")
        if "size" in kw:
            rr.font.size = kw["size"]
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.4), Inches(1.8), Inches(3.8), Inches(0.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0xE7, 0x4C, 0x3C)
    tb = s.shapes.add_textbox(Inches(5.4), Inches(2.8), Inches(3.8), Inches(1.5))
    tb.text_frame.word_wrap = True
    rr = tb.text_frame.paragraphs[0].add_run()
    rr.text = "Right-aligned red text"
    rr.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    rr.font.size = Pt(24)
    tb.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    # 8. Edge check: bars flush with all four slide edges, text on the bottom
    # edge - verifies nothing is clipped or hidden by the stage strip.
    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_title(s, "Edges: all four bars must be fully visible")
    W, H, T = prs.slide_width, prs.slide_height, Inches(0.25)
    for left, top, width, height, rgb in [
            (0, 0, W, T, (0xE7, 0x4C, 0x3C)), (0, H - T, W, T, (0x27, 0xAE, 0x60)),
            (0, 0, T, H, (0x29, 0x80, 0xB9)), (W - T, 0, T, H, (0x8E, 0x44, 0xAD))]:
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(*rgb)
        bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.5), H - T - Inches(0.6), Inches(9), Inches(0.5))
    tb.text_frame.text = "This line sits just above the bottom (green) bar"

    prs.save(args.out)
    print("wrote", args.out, "-", len(prs.slides.__iter__.__self__._sldIdLst), "slides")


if __name__ == "__main__":
    main()
